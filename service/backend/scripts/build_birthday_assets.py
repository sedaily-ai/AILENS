"""내 생일 타임머신 '스무 살의 내가 100만 원을 어디에 뒀다면' 공식 통계 수집 — 연도별 정적 JSON 과 최신값 JSON 을 만든다.

결과: service/frontend/public/data/birthday-assets/{YYYY}.json, latest.json
- 기준일은 서른 살 생일이 지났으면 스무 살 생일, 아직이면 태어난 날이다. 연도 파일 YYYY 는 기준일의 연도다.
  웹은 기준일 연도 파일과 latest.json 두 개만 받아 기기에서 계산한다. 서버에는 연도만 드러나고 생일 날짜는 가지 않는다.
- 출처는 한국은행 ECOS 네 계열이다: 소비자물가(901Y009, 국가데이터처), 정기예금 신규취급액 금리(121Y002),
  원/달러 종가 일별(731Y003)·매매기준율 월평균(731Y004), 서울 아파트 실거래가격지수(901Y089, 한국부동산원).
  금 한 돈은 세계은행 Pink Sheet 월별 엑셀의 Gold(CC BY 4.0)를 웹이 731Y004 월평균 환율로 바꾼다(--pink-sheet 필요).
  이용 근거가 확인된 출처만 수집한다(SOURCES 의 allowed). 근거는 docs/product/birthday-assets/출처_이용조건.md.
  허용된 출처가 하나도 없으면 파일을 쓰지 않는다.
- 정기예금은 금리를 매달 그달 금리로 굴린 누적지수(세전 월복리, deposit_index)로 바꿔 넣는다.
- 일별 계열은 연도 파일마다 전년 마지막 TAIL_DAYS 영업일을 함께 넣는다(1월 1일·연초 휴일 기준일의 직전 영업일).
- 연도 파일에는 생성 시각을 넣지 않는다(매달 갱신 때 지난 해 파일이 바뀌지 않게). 생성 시각은 latest.json 에만.

실행: cd service/backend && AWS_DEFAULT_REGION=us-east-1 uv run --no-project --with boto3==1.35.0 \\
        --with requests==2.32.3 --with openpyxl==3.1.5 python scripts/build_birthday_assets.py \\
        [--until YYYY-MM-DD] --pink-sheet <세계은행 CMO-Historical-Data-Monthly.xlsx 경로 또는 URL>
  Pink Sheet 주소(문서 ID)는 해가 바뀌면 달라진다 — 실행 전에
  https://www.worldbank.org/en/research/commodity-markets 의 'Monthly prices' 링크를 확인해 쓴다.
ECOS 인증키는 SSM(/sedaily-mbti/ecos-api-key)에서만 읽고 출력·기록하지 않는다.
"""
import argparse
import io
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import requests

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from common.dates.validation import KST  # noqa: E402
from common.secrets import get_secret  # noqa: E402

ECOS_KEY_PARAM = '/sedaily-mbti/ecos-api-key'
ECOS_URL = 'https://ecos.bok.or.kr/api/StatisticSearch/{key}/json/kr/{first}/{last}/{table}/{cycle}/{start}/{end}/{items}'
ECOS_PAGE = 100000  # 2026-10-20 실측: 일별 17,529건도 한 번에 내려왔다.
ECOS_TIMEOUT_SECONDS = 60
_ECOS_CODE = re.compile(r'^[A-Z]+-\d{3}$')  # ECOS RESULT.CODE 형식(INFO-200, ERROR-300 등)
OUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', 'frontend', 'public', 'data', 'birthday-assets'))
FIRST_YEAR = 1964  # 원/달러 기록 시작(1964-05-04)
TAIL_DAYS = 10

Point = Tuple[str, float]


@dataclass(frozen=True)
class Source:
    label: str                      # 화면 출처 표기
    allowed: bool                   # 출처_이용조건.md 판정이 '가능'일 때만 True
    requires: Tuple[str, ...] = ()  # 계산에 함께 쓰는 다른 출처(금의 원화 환산 → 환율)


# 판정이 바뀌면 여기와 출처_이용조건.md 를 함께 고친다(2026-10-21 승인안).
SOURCES: Dict[str, Source] = {
    'cpi': Source('국가데이터처 소비자물가지수(한국은행 ECOS 901Y009)', True),
    'deposit': Source('한국은행 예금은행 정기예금 신규취급액 가중평균금리(한국은행 ECOS 121Y002)', True),
    'usdkrw': Source('한국은행 원/달러 환율 종가·매매기준율 월평균(한국은행 ECOS 731Y003·731Y004)', True),
    'apartment': Source('한국부동산원 서울 아파트 실거래가격지수(한국은행 ECOS 901Y089)', True),
    # 2026-10-21 켬(사용자 요청). 세계은행 Pink Sheet 는 CC BY 4.0 이라 상업 이용이 되고 출처 표시가 의무다
    # (화면 출처 줄에 '세계은행 Pink Sheet(CC BY 4.0)'). 원자료(Bloomberg·Kitco·IMF IFS·LBM 등)의 제3자 권리는
    # 세계은행이 보증하지 않아 법무 확인 대상이다 — 출시(웹 배포) 전에 확인한다.
    'gold': Source('세계은행 Pink Sheet 금 월평균(CC BY 4.0)', True, requires=('usdkrw',)),
}

# ECOS 계열: (통계표, 주기, 시작, 끝(None 이면 --until), 항목 코드 경로). 2026-10-21 ECOS 로 확인한 값.
ECOS_SERIES = {
    'cpi': ('901Y009', 'M', '196501', None, '0'),                            # 총지수(2020=100)
    'deposit_rate': ('121Y002', 'M', '199601', None, 'BEABAA211'),           # 정기예금 신규취급액 가중평균(연리 %)
    'usdkrw_daily': ('731Y003', 'D', '19900302', None, '0000003'),           # 원/달러 종가 15:30
    'usdkrw_monthly': ('731Y004', 'M', '196405', None, '0000001/0000100'),   # 0000100 = 평균자료
    'apartment': ('901Y089', 'M', '200601', None, '200'),                    # 서울 아파트 실거래가격지수(2026.06=100)
}
SERIES_OF = {
    'cpi': ('cpi',),
    'deposit': ('deposit_rate',),
    'usdkrw': ('usdkrw_daily', 'usdkrw_monthly'),
    'apartment': ('apartment',),
}


class EcosError(RuntimeError):
    """ECOS 요청이 실패했거나 오류 코드를 돌려줬다(인증키 문제 등). 메시지에 키를 담지 않는다."""


def allowed_items(sources: Optional[Dict[str, Source]] = None) -> List[str]:
    """허용된 출처 중 함께 쓰는 출처까지 허용된 것만."""
    sources = SOURCES if sources is None else sources
    return [name for name, source in sources.items()
            if source.allowed and all(sources[need].allowed for need in source.requires)]


def _period(raw: str, cycle: str) -> str:
    if cycle == 'D':
        return f'{raw[:4]}-{raw[4:6]}-{raw[6:8]}'
    if cycle == 'M':
        return f'{raw[:4]}-{raw[4:6]}'
    return raw[:4]


def _next_month(period: str) -> str:
    year, month = int(period[:4]), int(period[5:7])
    year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return f'{year:04d}-{month:02d}'


def deposit_index(rates: Sequence[Point]) -> Tuple[List[Point], Optional[Point]]:
    """월별 연리(%) → 월초 누적지수. I(첫 달)=1, I(m+1)=I(m)×(1+r_m/1200)(세전 월복리).
    반환: (금리가 있는 달의 지수 목록, 마지막 금리 다음 달의 지수). 빠진 달이 있으면 ValueError."""
    index: List[Point] = []
    value = 1.0
    previous: Optional[str] = None
    for period, rate in rates:
        if previous is not None and period != _next_month(previous):
            raise ValueError(f'정기예금 금리가 연속이 아니다: {previous} 다음이 {period}')
        index.append((period, round(value, 10)))
        value *= 1 + rate / 1200
        previous = period
    if previous is None:
        return [], None
    return index, (_next_month(previous), round(value, 10))


def fetch_ecos(table: str, cycle: str, start: str, end: str, items: str, key: str,
               get: Callable = requests.get) -> List[Point]:
    """ECOS StatisticSearch 한 계열. 자료가 없으면 []. 다른 오류는 EcosError(키 든 URL 은 메시지·트레이스백에 남기지 않는다)."""
    points: List[Point] = []
    first = 1
    while True:
        url = ECOS_URL.format(key=key, first=first, last=first + ECOS_PAGE - 1, table=table, cycle=cycle,
                              start=start, end=end, items=items)
        try:
            body = get(url, timeout=ECOS_TIMEOUT_SECONDS).json()
        except (requests.RequestException, ValueError) as error:
            raise EcosError(f'ECOS {table} 요청 실패: {type(error).__name__}') from None
        found = body.get('StatisticSearch')
        if not found:
            result = body.get('RESULT') or {}
            code = result.get('CODE')
            if code == 'INFO-200':  # 해당하는 데이터가 없음
                return points
            # MESSAGE 는 요청 값(인증키 포함)을 되풀이할 수 있어 넣지 않는다. CODE 도 ECOS 형식(INFO-100 등)일 때만 쓴다.
            shown = code if isinstance(code, str) and _ECOS_CODE.match(code) else '알 수 없는 코드'
            raise EcosError(f'ECOS {table} 오류: {shown}')
        rows = found.get('row') or []
        for row in rows:
            value = row.get('DATA_VALUE')
            if value not in (None, ''):
                points.append((_period(str(row['TIME']), cycle), float(value)))
        total = int(found.get('list_total_count') or 0)
        first += len(rows)
        if not rows or first > total:
            return sorted(points)


_PINK_MONTH = re.compile(r'^(\d{4})M(\d{2})$')


def parse_pink_sheet_gold(data: bytes) -> List[Point]:
    """세계은행 CMO Historical Data Monthly 엑셀의 'Monthly Prices' 시트에서 금(달러/트로이온스) 월평균.
    열 위치는 매달 바뀔 수 있어 열 제목이 'Gold' 인 열을 찾는다. 행은 '1960M01' 형식."""
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True)['Monthly Prices']
    rows = list(sheet.iter_rows(values_only=True))
    column = None
    for row in rows:
        for index, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip() == 'Gold':
                column = index
                break
        if column is not None:
            break
    if column is None:
        raise ValueError("'Monthly Prices' 시트에서 'Gold' 열을 찾지 못했다")
    points = []
    for row in rows:
        match = _PINK_MONTH.match(str(row[0]).strip()) if row and row[0] is not None else None
        if match and column < len(row) and isinstance(row[column], (int, float)):
            points.append((f'{match.group(1)}-{match.group(2)}', float(row[column])))
    return sorted(points)


def read_source(path_or_url: str, get: Callable = requests.get) -> bytes:
    if re.match(r'^https?://', path_or_url):
        res = get(path_or_url, timeout=ECOS_TIMEOUT_SECONDS)
        res.raise_for_status()
        return res.content
    with open(path_or_url, 'rb') as fp:
        return fp.read()


def _daily_slice(points: Sequence[Point], year: int) -> List[Point]:
    start = f'{year}-01-01'
    before = [p for p in points if p[0] < start][-TAIL_DAYS:]
    return before + [p for p in points if p[0].startswith(f'{year}-')]


def _year_slice(points: Sequence[Point], year: int) -> List[Point]:
    return [p for p in points if p[0].startswith(f'{year}-')]


def build_year_files(series: dict, items: Sequence[str], first_year: int, last_year: int) -> Dict[int, dict]:
    """연도 → 파일 내용. 허용된 항목의 키만 넣는다. 자료 시작 전 연도는 빈 배열."""
    files = {}
    for year in range(first_year, last_year + 1):
        data: dict = {'year': year}
        if 'cpi' in items:
            data['cpi'] = _year_slice(series['cpi'], year)
        if 'deposit' in items:
            data['deposit'] = _year_slice(series['deposit_index'], year)
        if 'usdkrw' in items:
            data['usdkrw'] = {'daily': _daily_slice(series['usdkrw_daily'], year),
                              'monthly': _year_slice(series['usdkrw_monthly'], year)}
        if 'apartment' in items:
            data['apartment'] = _year_slice(series['apartment'], year)
        if 'gold' in items:
            data['gold'] = _year_slice(series['gold'], year)
        files[year] = data
    return files


def build_latest(series: dict, items: Sequence[str], generated_at: str, until: str,
                 sources: Optional[Dict[str, Source]] = None) -> dict:
    """최신값. last 는 ECOS 원계열별 마지막 시점(예금은 마지막 금리 달), deposit 은 그 다음 달 누적지수."""
    sources = SOURCES if sources is None else sources
    point = lambda name: list(series[name][-1]) if series.get(name) else None  # noqa: E731
    period = lambda name: series[name][-1][0] if series.get(name) else None  # noqa: E731
    last: dict = {}
    latest: dict = {'generatedAt': generated_at, 'until': until,
                    'sources': {name: sources[name].label for name in items}, 'last': last}
    if 'cpi' in items:
        last['cpi'] = period('cpi')
        latest['cpi'] = point('cpi')
    if 'deposit' in items:
        last['deposit_rate'] = period('deposit_rate')
        latest['deposit'] = list(series['deposit_next']) if series.get('deposit_next') else None
    if 'usdkrw' in items:
        last['usdkrw_daily'] = period('usdkrw_daily')
        last['usdkrw_monthly'] = period('usdkrw_monthly')
        latest['usdkrw'] = {'daily': point('usdkrw_daily'), 'monthly': point('usdkrw_monthly')}
    if 'apartment' in items:
        last['apartment'] = period('apartment')
        latest['apartment'] = point('apartment')
    if 'gold' in items:
        last['gold'] = period('gold')
        fx = dict(series.get('usdkrw_monthly', []))
        common = [p for p in series.get('gold', []) if p[0] in fx]  # 금·환율 월평균이 둘 다 있는 마지막 달
        latest['gold'] = ({'usd': list(common[-1]), 'usdkrw': [common[-1][0], fx[common[-1][0]]]}
                          if common else None)
    return latest


def collect(items: Sequence[str], until: str, key: Optional[str], pink_sheet: Optional[str],
            get: Callable = requests.get) -> dict:
    """허용된 항목에 필요한 계열만 받는다. 정기예금은 누적지수(deposit_index·deposit_next)도 함께 둔다."""
    series: dict = {}
    for item in items:
        for name in SERIES_OF.get(item, ()):
            table, cycle, start, end, codes = ECOS_SERIES[name]
            if end is None:
                end = until.replace('-', '') if cycle == 'D' else until[:7].replace('-', '')
            series[name] = fetch_ecos(table, cycle, start, end, codes, key, get=get)
        if item == 'deposit':
            series['deposit_index'], series['deposit_next'] = deposit_index(series['deposit_rate'])
    if 'gold' in items:
        if not pink_sheet:
            raise SystemExit('금을 수집하려면 --pink-sheet 로 세계은행 월별 엑셀 경로나 URL 을 넘겨야 한다')
        series['gold'] = [p for p in parse_pink_sheet_gold(read_source(pink_sheet, get=get)) if p[0] <= until[:7]]
    return series


def write_files(out_dir: str, files: Dict[int, dict], latest: dict) -> List[int]:
    """파일을 쓰고 각 파일 크기(바이트)를 돌려준다."""
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for year, data in files.items():
        path = os.path.join(out_dir, f'{year}.json')
        with open(path, 'w', encoding='utf-8') as fp:
            json.dump(data, fp, ensure_ascii=False, separators=(',', ':'))
        paths.append(path)
    path = os.path.join(out_dir, 'latest.json')
    with open(path, 'w', encoding='utf-8') as fp:
        json.dump(latest, fp, ensure_ascii=False, indent=1)
        fp.write('\n')
    paths.append(path)
    return [os.path.getsize(p) for p in paths]


def main(argv=None, get: Callable = requests.get, secret: Callable[[str], str] = get_secret) -> int:
    now = datetime.now(KST)
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--until', default=(now - timedelta(days=1)).strftime('%Y-%m-%d'))
    parser.add_argument('--out', default=OUT_DIR)
    parser.add_argument('--pink-sheet', default=None)
    args = parser.parse_args(argv)

    items = allowed_items()
    if not items:
        print('이용 근거가 확인된 출처가 없어 아무 파일도 쓰지 않았다(SOURCES 의 allowed, 출처_이용조건.md).')
        return 0
    needs_ecos = any(item in SERIES_OF for item in items)
    key = secret(ECOS_KEY_PARAM) if needs_ecos else None
    series = collect(items, args.until, key, args.pink_sheet, get=get)
    last_year = max(int(args.until[:4]), now.year)  # 1월 초 실행에도 올해 파일(전년 꼬리)이 있게
    files = build_year_files(series, items, FIRST_YEAR, last_year)
    sizes = write_files(args.out, files, build_latest(series, items, now.isoformat(timespec='seconds'), args.until))
    raw = [name for item in items for name in SERIES_OF.get(item, ())] + (['gold'] if 'gold' in items else [])
    summary = ', '.join(f'{name} {len(series[name])}건({series[name][0][0]}~{series[name][-1][0]})'
                        if series[name] else f'{name} 0건' for name in raw)
    print(f'items={",".join(items)} years={FIRST_YEAR}~{last_year} {summary} '
          f'files={len(sizes)} total={sum(sizes) / 1024:.1f}KB max={max(sizes) / 1024:.1f}KB -> {args.out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
